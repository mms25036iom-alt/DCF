"""Portable XLSX export for the app's stock-specific DCF.

Uses the app's existing openpyxl runtime dependency. Excel formulas remain
editable, with baseline calculation caches for readers that do not recalculate.
"""

from __future__ import annotations

import io
import math
import re
import zipfile
from datetime import date
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

from research.data.stock_model import ASSUMPTION_BOUNDS, HISTORY_METRICS, finite_number
from research.reporting.live_report import public_source_url, redact_private_locations

NAVY = "172E36"
TEAL = "18786B"
MUTED = "6F807A"
PALE = "EDF4EF"
BLUE = "1768B4"
GREEN = "18786B"
MONEY = '#,##0.00;[Red](#,##0.00);"–"'
PERCENT = '0.0%;[Red](0.0%);"–"'
INTEGER = '#,##0;[Red](#,##0);"–"'
METRIC_LABELS = {
    "revenue": "Revenue",
    "ebit": "Operating profit (EBIT)",
    "net_income": "Net income",
    "ebitda": "EBITDA",
    "operating_cash_flow": "Operating cash flow",
    "capital_expenditure": "Capital expenditure",
    "depreciation": "Depreciation & amortization",
    "total_debt": "Total debt",
    "cash_and_equivalents": "Cash & equivalents",
    "shareholders_equity": "Shareholders' equity",
    "income_tax": "Income tax",
    "pretax_income": "Pretax income",
    "working_capital_cash_flow": "Working-capital cash flow",
    "accounts_receivable": "Accounts receivable",
    "inventory": "Inventory",
    "accounts_payable": "Accounts payable",
    "cost_of_revenue": "Cost of revenue",
    "current_assets": "Current assets",
    "current_liabilities": "Current liabilities",
    "interest_expense": "Interest expense",
    "noncontrolling_interest": "Non-controlling interests",
}
INPUT_ROWS = {
    "revenue_growth": 11,
    "ebit_margin": 12,
    "tax_rate": 13,
    "depreciation_rate": 14,
    "capex_rate": 15,
    "working_capital_rate": 16,
    "discount_rate": 17,
    "terminal_growth": 18,
    "net_debt": 19,
    "shares_outstanding": 20,
    "risk_free_rate": 24,
    "equity_risk_premium": 25,
    "beta": 26,
    "pretax_cost_of_debt": 27,
    "debt_weight": 28,
    "wacc_override": 29,
    "working_capital_method": 33,
    "receivable_days": 34,
    "inventory_days": 35,
    "payable_days": 36,
    "cogs_rate": 37,
    "revenue_growth_end": 39,
    "terminal_roic": 40,
}
INPUT_LABELS = {
    "revenue_growth": "First forecast revenue growth",
    "revenue_growth_end": "Fifth forecast revenue growth",
    "terminal_roic": "Terminal return on invested capital",
    "ebit_margin": "EBIT margin",
    "tax_rate": "Tax rate",
    "depreciation_rate": "Depreciation / revenue",
    "capex_rate": "Capital expenditure / revenue",
    "working_capital_rate": "Operating working capital / revenue",
    "discount_rate": "Manual WACC (manual mode only)",
    "terminal_growth": "Terminal growth",
    "net_debt": "Net debt / (net cash)",
    "shares_outstanding": "Shares outstanding (absolute)",
    "risk_free_rate": "Risk-free rate assumption",
    "equity_risk_premium": "Equity risk premium assumption",
    "beta": "Levered beta",
    "pretax_cost_of_debt": "Pre-tax borrowing cost",
    "debt_weight": "Debt / total capital weight",
    "wacc_override": "WACC mode: 0 CAPM / 1 manual",
    "working_capital_method": "Working-capital mode: 0 ratio / 1 days",
    "receivable_days": "Receivable days (DSO)",
    "inventory_days": "Inventory days (DIO)",
    "payable_days": "Payable days (DPO)",
    "cogs_rate": "Cost of revenue / revenue",
}
YEARLY_INPUT_ROWS = {
    "revenue_growth": 46,
    "ebit_margin": 47,
    "tax_rate": 48,
    "depreciation_rate": 49,
    "capex_rate": 50,
    "working_capital_rate": 51,
    "receivable_days": 52,
    "inventory_days": 53,
    "payable_days": 54,
    "cogs_rate": 55,
}
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _safe_text(value) -> str:
    # User/provider strings are always explicit strings, never executable cells.
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", redact_private_locations(str(value or "")))[
        :32760
    ]


def _iso_date(value):
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def _write_text(sheet, row: int, column: int, value):
    cell = sheet.cell(row, column, _safe_text(value))
    cell.data_type = "s"
    return cell


def _title(sheet, title: str, subtitle: str):
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:H2")
    _write_text(sheet, 1, 1, title).font = Font(
        name="Aptos Display", size=22, bold=True, color="FFFFFF"
    )
    sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["A1"].alignment = Alignment(vertical="center")
    for row in sheet["A1:H2"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=NAVY)
    sheet.row_dimensions[1].height = 25
    sheet.row_dimensions[2].height = 17
    sheet.merge_cells("A3:H3")
    _write_text(sheet, 3, 1, subtitle).font = Font(name="Aptos", size=10, color=MUTED)
    sheet.row_dimensions[3].height = 22
    for column in range(1, 9):
        sheet.column_dimensions[get_column_letter(column)].width = 19 if column > 1 else 37
    sheet.freeze_panes = "B7"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = "Astra Research | Assumption-driven scenario"
    sheet.oddFooter.right.text = "Page &P of &N"
    sheet.sheet_properties.tabColor = TEAL


def _section(sheet, row: int, title: str, end_column=8):
    for column in range(1, end_column + 1):
        sheet.cell(row, column).fill = PatternFill("solid", fgColor=PALE)
        sheet.cell(row, column).font = Font(name="Aptos", bold=True, size=10, color=TEAL)
    _write_text(sheet, row, 1, title)
    sheet.row_dimensions[row].height = 23


def _formula(sheet, coordinate: str, expression: str, value, caches: dict, fmt=MONEY):
    cell = sheet[coordinate]
    cell.value = expression
    cell.font = Font(name="Aptos", size=11, color=GREEN if "!" in expression else NAVY)
    cell.number_format = fmt
    caches.setdefault(sheet.title, {})[coordinate] = value
    return cell


def _note(sheet, row: int, text: str, end_column=8):
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_column)
    cell = _write_text(sheet, row, 1, text)
    cell.font = Font(name="Aptos", size=10, color=MUTED)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[row].height = max(26, 14 * math.ceil(len(text) / 130) + 8)


def _cached_bytes(workbook: Workbook, caches: dict) -> bytes:
    stream = io.BytesIO()
    workbook.save(stream)
    output = io.BytesIO()
    # openpyxl cannot write cached formula values. Populate only values already
    # calculated by the same model; formulas themselves remain the source of
    # future Excel recalculation. No external links or macros are introduced.
    with (
        zipfile.ZipFile(stream) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for item in source.infolist():
            data = source.read(item.filename)
            match = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", item.filename)
            if match:
                sheet_name = workbook.sheetnames[int(match[1]) - 1]
                values = caches.get(sheet_name, {})
                if values:
                    root = ET.fromstring(data)
                    for cell in root.findall(f".//{{{NS}}}c"):
                        coordinate = cell.attrib.get("r")
                        if coordinate not in values or cell.find(f"{{{NS}}}f") is None:
                            continue
                        value = values[coordinate]
                        cached = cell.find(f"{{{NS}}}v")
                        if cached is None:
                            cached = ET.SubElement(cell, f"{{{NS}}}v")
                        if value is None:
                            cached.text = None
                        elif isinstance(value, str):
                            cell.set("t", "str")
                            cached.text = value
                        elif isinstance(value, (int, float)) and math.isfinite(value):
                            cell.attrib.pop("t", None)
                            cached.text = repr(value)
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(item, data)
    return output.getvalue()


def build_stock_workbook(model: dict) -> bytes:
    """Export one linked operating model, with original observations kept separate."""
    wb = Workbook()
    wb.remove(wb.active)
    names = (
        "Summary",
        "Assumptions",
        "Historical",
        "OperatingForecast",
        "WorkingCapital",
        "WACC",
        "DCF",
        "Sensitivity",
        "Sources",
    )
    for name in names:
        wb.create_sheet(name)
    wb.calculation = CalcProperties(
        calcId=191029, fullCalcOnLoad=True, forceFullCalc=True, calcMode="auto"
    )
    wb.properties.creator = "Astra Research"
    wb.properties.title = _safe_text(f"{model.get('ticker')} | Linked equity valuation")
    caches: dict = {}
    unit = model["unit"]
    divisor = unit["divisor"]
    history = model.get("historical", [])
    snapshot = model.get("source_snapshot", {})
    snapshot = {
        **snapshot,
        **{
            key: finite_number(snapshot.get(key))
            for key in ("shares_outstanding", "current_price", "beta", "market_cap")
        },
    }
    inputs = model.get("assumptions", {})
    notes = model.get("assumption_notes", {})
    forecast = model.get("forecast", [])
    valuation = model.get("valuation") or {}
    company = model.get("company") or model.get("ticker") or "Requested listing"
    for name in names:
        _title(
            wb[name],
            f"{company} | {name}",
            f"{model.get('ticker')} | "
            + (
                f"{model.get('currency')} absolute source amounts; shares / prices as stated"
                if name == "Sources"
                else f"{unit['amount_label']} unless stated | Assumption-driven valuation"
            ),
        )

    def text(sheet, row, col, value):
        return _write_text(wb[sheet], row, col, value)

    def number(sheet, row, col, value, fmt=MONEY):
        cell = wb[sheet].cell(row, col, finite_number(value))
        cell.number_format = fmt
        return cell

    def formula(sheet, row, col, expression, value, fmt=MONEY):
        return _formula(wb[sheet], f"{get_column_letter(col)}{row}", expression, value, caches, fmt)

    def safe_url(value):
        return public_source_url(value) or "Reference withheld: not a public source URL"

    def date_cell(sheet, row, column, value):
        observed = _iso_date(value)
        if observed:
            cell = wb[sheet].cell(row, column, to_excel(observed))
            cell.number_format = "yyyy-mm-dd"
            return cell

    source = wb["Sources"]
    for row, label, value in (
        (4, "Company", company),
        (5, "Ticker", model.get("ticker")),
        (6, "Currency", model.get("currency")),
    ):
        text("Sources", row, 1, label)
        text("Sources", row, 2, value)
    for row, label, value in (
        (7, "Current shares (absolute)", snapshot.get("shares_outstanding")),
        (8, "Observed price / share", snapshot.get("current_price")),
        (9, "Reported beta", snapshot.get("beta")),
        (10, "Market capitalization (absolute)", snapshot.get("market_cap")),
    ):
        text("Sources", row, 1, label)
        number("Sources", row, 2, value, INTEGER if row == 7 else MONEY)
    text("Sources", 11, 1, "Retrieved at")
    text("Sources", 11, 2, snapshot.get("retrieved_at"))
    text("Sources", 12, 1, "Quote as of")
    text("Sources", 12, 2, snapshot.get("price_as_of"))
    source_rows = {metric: 14 + i for i, metric in enumerate(HISTORY_METRICS)}
    for col, record in enumerate(history, 2):
        text("Sources", 13, col, record.get("period", record.get("year")))
        text("Historical", 6, col, record.get("period", record.get("year")))
        for index, metric in enumerate(HISTORY_METRICS):
            row = source_rows[metric]
            text("Sources", row, 1, METRIC_LABELS.get(metric, metric))
            text("Historical", index + 7, 1, METRIC_LABELS.get(metric, metric))
            raw = record.get(metric)
            number("Sources", row, col, raw)
            if raw is not None:
                formula(
                    "Historical",
                    index + 7,
                    col,
                    f"=Sources!{get_column_letter(col)}{row}/Assumptions!$B$6",
                    raw / divisor,
                )
    fiscal_date_row = max(source_rows.values(), default=34) + 1
    valuation_date_row = fiscal_date_row + 1
    text("Sources", fiscal_date_row, 1, "Fiscal period end")
    for col, record in enumerate(history, 2):
        date_cell("Sources", fiscal_date_row, col, record.get("period"))
    text("Sources", valuation_date_row, 1, "Valuation date")
    date_cell("Sources", valuation_date_row, 2, model.get("valuation_date"))
    citation_row = max(source_rows.values(), default=34) + 4
    _section(source, citation_row, "Source references and limitations")
    for offset, citation in enumerate(snapshot.get("citations", []), 1):
        row = citation_row + offset
        source.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        source.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
        title = text("Sources", row, 1, citation.get("title", "Source"))
        link = text("Sources", row, 3, safe_url(citation.get("uri")))
        title.alignment = link.alignment = Alignment(wrap_text=True, vertical="center")
        source.row_dimensions[row].height = 32
    _note(
        source,
        citation_row + len(snapshot.get("citations", [])) + 3,
        "Original currency observations. Blank cells are unavailable; no FX conversion or missing-value fill is applied.",
    )
    macro_row = citation_row + len(snapshot.get("citations", [])) + 6
    _section(source, macro_row, "Dated market assumptions and source provenance")
    macro_rows = {}
    for key, observation in (snapshot.get("market_assumptions") or {}).items():
        if not isinstance(observation, dict):
            continue
        macro_row += 1
        macro_rows[key] = macro_row
        text("Sources", macro_row, 1, key.replace("_", " ").title())
        number("Sources", macro_row, 2, observation.get("value"), PERCENT)
        text("Sources", macro_row, 3, observation.get("as_of"))
        source.merge_cells(start_row=macro_row, start_column=4, end_row=macro_row, end_column=8)
        text(
            "Sources", macro_row, 4, safe_url(observation.get("source_url"))
        ).alignment = Alignment(wrap_text=True)
        source.row_dimensions[macro_row].height = 32
    _note(
        source,
        macro_row + 2,
        "Beta is a provider estimate when available; borrowing cost and capital weights use disclosed book-debt proxies. Reported non-controlling interests are a book-value proxy for a market-value claim. Review observation dates and methodology.",
    )

    summary = wb["Summary"]
    _section(summary, 5, "Valuation overview")
    for row, label in (
        (8, "Enterprise value"),
        (9, "Equity value"),
        (12, "Illustrative value / share"),
        (13, "Observed market price / share"),
        (14, "Model / market difference"),
        (16, "Active WACC"),
        (17, "Terminal growth"),
    ):
        text("Summary", row, 1, label)
    _note(
        summary,
        20,
        "Forecasts are editable scenarios, not audited observations or price targets. Blue cells are editable inputs; green cells link schedules. All forecasts use one active set of assumptions.",
    )
    _note(
        summary,
        23,
        "Change Assumptions or source observations to recalculate the whole workbook. Review source dates, missing facts, and industry suitability before interpreting results.",
    )
    quality = model.get("quality") or {}
    _note(
        summary,
        26,
        f"Input review: {quality.get('status', 'needs_review')}. Sourced inputs have not been independently verified or audited.",
    )
    for row, issue in enumerate(quality.get("issues") or [], 28):
        _note(summary, row, str(issue))
    if not inputs or not forecast:
        reason = (
            model.get("withheld_reason")
            or "A supported DCF cannot be built from these observations."
        )
        for name in ("DCF", "WorkingCapital", "WACC", "OperatingForecast", "Sensitivity"):
            _note(wb[name], 6, reason)
        _note(summary, 6, reason)
        return _cached_bytes(wb, caches)

    base_index = next(
        i for i, record in enumerate(history) if record.get("period") == model.get("base_period")
    )
    base = history[base_index]
    base_col = get_column_letter(base_index + 2)

    def ref(metric):
        return f"Sources!${base_col}${source_rows[metric]}"

    assumptions = wb["Assumptions"]
    text("Assumptions", 6, 1, "Monetary display divisor")
    number("Assumptions", 6, 2, divisor, INTEGER)
    text("Assumptions", 7, 1, "Base revenue")
    formula("Assumptions", 7, 2, f"={ref('revenue')}/B6", base["revenue"] / divisor)
    text("Assumptions", 8, 1, "Base financial year")
    number("Assumptions", 8, 2, model["base_year"], "0")
    text("Assumptions", 9, 1, "Base observations")
    text("Assumptions", 9, 2, model.get("base_period"))
    calibrated = {
        "net_debt": f"=({ref('total_debt')}-{ref('cash_and_equivalents')})/$B$6",
        "shares_outstanding": "=Sources!B7",
        "revenue_growth_end": "=$B$18",
        "terminal_roic": "=WACC!B20",
    }
    normalization = model.get("normalization") or {}
    period_columns = {
        str(row.get("period")): get_column_letter(i + 2) for i, row in enumerate(history)
    }

    def period_ref(period, metric):
        return f"Sources!${period_columns[str(period)]}${source_rows[metric]}"

    growth_periods = normalization.get("revenue_growth", [])
    if len(growth_periods) == 2 and all(str(p) in period_columns for p in growth_periods):
        earlier, later = [
            next(r for r in history if str(r.get("period")) == str(p)) for p in growth_periods
        ]
        calibrated["revenue_growth"] = (
            f"=({period_ref(growth_periods[-1], 'revenue')}/{period_ref(growth_periods[0], 'revenue')})^(1/({later['year']}-{earlier['year']}))-1"
        )
    ratio_metrics = {
        "ebit_margin": ("ebit", "revenue", 1),
        "depreciation_rate": ("depreciation", "revenue", 1),
        "capex_rate": ("capital_expenditure", "revenue", 1),
        "tax_rate": ("income_tax", "pretax_income", 1),
        "receivable_days": ("accounts_receivable", "revenue", 365),
        "inventory_days": ("inventory", "cost_of_revenue", 365),
        "payable_days": ("accounts_payable", "cost_of_revenue", 365),
        "cogs_rate": ("cost_of_revenue", "revenue", 1),
    }
    for key, (numerator, denominator, scale) in ratio_metrics.items():
        terms = []
        for period in normalization.get(key, []):
            if str(period) not in period_columns:
                continue
            top, bottom = period_ref(period, numerator), period_ref(period, denominator)
            value = f"ABS({top})" if key == "capex_rate" else top
            condition = f"AND(ISNUMBER({top}),ISNUMBER({bottom}),{bottom}>0)"
            if key not in {"ebit_margin", "capex_rate"}:
                condition = f"AND({condition},{top}>=0)"
            if key == "tax_rate":
                condition = f"AND({condition},{top}/{bottom}<=0.6)"
            terms.append(f'IF({condition},{value}/{bottom}*{scale},"")')
        if terms:
            calibrated[key] = f'=IFERROR(MEDIAN({",".join(terms)}),"")'
    market = snapshot.get("market_assumptions") or {}

    def uses_market_input(key):
        return (
            key in macro_rows
            and finite_number((market.get(key) or {}).get("value")) == inputs.get(key)
            and not notes.get(key, "").startswith("Explicit user")
        )

    for key in ("risk_free_rate", "equity_risk_premium"):
        if uses_market_input(key):
            calibrated[key] = f"=Sources!B{macro_rows[key]}"
    if (
        uses_market_input("risk_free_rate")
        and (market.get("risk_free_rate") or {}).get("adjustment_source_url")
        and "sovereign_yield" in macro_rows
        and "sovereign_default_spread" in macro_rows
    ):
        calibrated["risk_free_rate"] = (
            f"=Sources!B{macro_rows['sovereign_yield']}-Sources!B{macro_rows['sovereign_default_spread']}"
        )
    if snapshot.get("beta") is not None and 0 <= snapshot["beta"] <= 5:
        calibrated["beta"] = "=Sources!B9"
    if (
        base.get("interest_expense") is not None
        and base.get("total_debt", 0) > 0
        and 0 <= abs(base["interest_expense"]) / base["total_debt"] <= 0.6
    ):
        calibrated["pretax_cost_of_debt"] = f"=ABS({ref('interest_expense')})/{ref('total_debt')}"
    cap_ref = (
        "Sources!$B$10" if (snapshot.get("market_cap") or 0) > 0 else "(Sources!$B$7*Sources!$B$8)"
    )
    if (
        snapshot.get("market_cap")
        or (snapshot.get("current_price") or 0) * (snapshot.get("shares_outstanding") or 0)
    ) > 0:
        calibrated["debt_weight"] = f"={ref('total_debt')}/({ref('total_debt')}+{cap_ref})"
    wc = model.get("working_capital", {})
    wc_terms = []
    for period in normalization.get("working_capital_rate", []):
        if str(period) not in period_columns:
            continue
        ar, inv, ap, sales = [
            period_ref(period, metric)
            for metric in ("accounts_receivable", "inventory", "accounts_payable", "revenue")
        ]
        wc_terms.append(
            f'IF(AND(ISNUMBER({ar}),ISNUMBER({inv}),ISNUMBER({ap}),ISNUMBER({sales}),{ar}>=0,{inv}>=0,{ap}>=0,{sales}>0),({ar}+{inv}-{ap})/{sales},"")'
        )
    if wc_terms:
        calibrated["working_capital_rate"] = f'=IFERROR(MEDIAN({",".join(wc_terms)}),"")'
    validations = {}
    for key, row in INPUT_ROWS.items():
        if key not in inputs:
            continue
        text("Assumptions", row, 1, INPUT_LABELS.get(key, key.replace("_", " ").title()))
        value = inputs[key] / divisor if key == "net_debt" else inputs[key]
        fmt = (
            "0"
            if key in {"working_capital_method", "wacc_override"}
            else INTEGER
            if key == "shares_outstanding"
            else MONEY
            if key in {"net_debt", "beta", "receivable_days", "inventory_days", "payable_days"}
            else PERCENT
        )
        if key in calibrated and not notes.get(key, "").startswith("Explicit user"):
            formula("Assumptions", row, 2, calibrated[key], value, fmt)
        else:
            cell = number("Assumptions", row, 2, value, fmt)
            cell.font = Font(name="Aptos", color=BLUE, size=11)
            cell.fill = PatternFill("solid", fgColor="F1F6FF")
        assumptions.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
        cell = text("Assumptions", row, 3, notes.get(key, "Explicit modeling input"))
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        assumptions.row_dimensions[row].height = 42
        low, high = ASSUMPTION_BOUNDS[key]
        if key == "net_debt":
            low, high = low / divisor, high / divisor
        validation = DataValidation(
            type="whole" if key in {"working_capital_method", "wacc_override"} else "decimal",
            operator="between",
            formula1=str(low),
            formula2=str(high),
        )
        validation.errorTitle = "Input outside allowed range"
        validation.error = "Use a valid modeling input; keep WACC greater than terminal growth."
        validation.showErrorMessage = True
        assumptions.add_data_validation(validation)
        validation.add(f"B{row}")
        validations[key] = validation
    _note(
        assumptions,
        42,
        "Modes: WACC 0 = linked CAPM, 1 = manual. Working capital 0 = revenue ratio, 1 = operating days. Rf/ERP/default borrowing costs are disclosed assumptions where no independent market observation is supplied.",
    )
    _section(assumptions, 43, "Year-by-year operating assumptions")
    _note(
        assumptions,
        44,
        "Yearly blue cells link to global drivers. Growth fades linearly from year one to year five. Replace a formula with an annual override when needed; retain formulas to inherit global changes.",
    )
    for col in range(3, 8):
        text("Assumptions", 45, col, f"{model['base_year'] + col - 2}E")
    for key, row in YEARLY_INPUT_ROWS.items():
        text("Assumptions", row, 1, INPUT_LABELS[key])
        for col in range(3, 8):
            cell = formula(
                "Assumptions",
                row,
                col,
                f"=$B$11+($B$39-$B$11)*({col}-3)/4"
                if key == "revenue_growth"
                else f"=$B${INPUT_ROWS[key]}",
                forecast[col - 3].get("revenue_growth", inputs[key])
                if key == "revenue_growth"
                else inputs[key],
                MONEY if key.endswith("_days") else PERCENT,
            )
            cell.font = Font(name="Aptos", color=BLUE, size=11)
            cell.fill = PatternFill("solid", fgColor="F1F6FF")
            validations[key].add(cell)

    _section(assumptions, 57, "Valuation timing and terminal policy")
    text("Assumptions", 58, 1, "Valuation date")
    formula(
        "Assumptions",
        58,
        2,
        f"=Sources!B{valuation_date_row}",
        to_excel(_iso_date(model["valuation_date"])),
        "yyyy-mm-dd",
    )
    text("Assumptions", 59, 1, "Base fiscal period end")
    formula(
        "Assumptions",
        59,
        2,
        f"=Sources!{base_col}{fiscal_date_row}",
        to_excel(_iso_date(base["period"])),
        "yyyy-mm-dd",
    )
    text("Assumptions", 60, 1, "Terminal ROIC mode: 0 WACC / 1 manual")
    explicit_roic = (model.get("terminal") or {}).get("method") == "explicit_roic"
    number("Assumptions", 60, 2, int(explicit_roic), "0").font = Font(name="Aptos", color=BLUE)
    roic_mode = DataValidation(type="whole", operator="between", formula1="0", formula2="1")
    roic_mode.showErrorMessage = True
    assumptions.add_data_validation(roic_mode)
    roic_mode.add("B60")
    text("Assumptions", 61, 1, "Annual coverage current (1 yes / 0 stale)")
    formula("Assumptions", 61, 2, "=IF(EDATE(B59,12)>B58,1,0)", 1, "0")
    text("Assumptions", 62, 1, "Dated quote comparison available (1 / 0)")
    comparison_available = (model.get("timing") or {}).get("market_comparison_available", True)
    number("Assumptions", 62, 2, int(comparison_available), "0")
    quote_mode = DataValidation(type="whole", operator="between", formula1="0", formula2="1")
    assumptions.add_data_validation(quote_mode)
    quote_mode.add("B62")
    _note(
        assumptions,
        63,
        "Discount periods use actual days / 365.25. First-year FCFF is multiplied by the remaining fraction of its fiscal year, assuming uniform cash generation. Later years retain full FCFF. Cash taxes apply only to positive EBIT.",
    )
    _note(
        assumptions,
        65,
        "Terminal ROIC defaults to active WACC (no excess returns). To use a fixed ROIC, set B60 to 1 and enter it in B40. Terminal reinvestment rate = terminal growth / ROIC. Sensitivity recomputes this rate for every WACC and growth combination.",
    )
    _section(assumptions, 68, "Historical calibration periods")
    for row, (key, periods) in enumerate(normalization.items(), 69):
        text("Assumptions", row, 1, INPUT_LABELS.get(key, key))
        assumptions.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        text("Assumptions", row, 2, ", ".join(str(p) for p in periods)).alignment = Alignment(
            wrap_text=True
        )

    wacc = model.get("wacc", {})
    wacc_links = {
        6: ("Risk-free rate", "=Assumptions!B24", inputs["risk_free_rate"]),
        7: ("Equity risk premium", "=Assumptions!B25", inputs["equity_risk_premium"]),
        8: ("Levered beta", "=Assumptions!B26", inputs["beta"]),
        10: ("Cost of equity (CAPM)", "=B6+B7*B8", wacc.get("cost_of_equity")),
        12: ("Pre-tax borrowing cost", "=Assumptions!B27", inputs["pretax_cost_of_debt"]),
        13: ("Tax rate", "=Assumptions!B13", inputs["tax_rate"]),
        14: ("After-tax cost of debt", "=B12*(1-B13)", wacc.get("after_tax_cost_of_debt")),
        16: ("Debt weight", "=Assumptions!B28", inputs["debt_weight"]),
        17: ("Equity weight", "=1-B16", 1 - inputs["debt_weight"]),
        19: ("Calculated WACC", "=B10*B17+B14*B16", wacc.get("calculated_wacc")),
        20: ("Active WACC", "=IF(Assumptions!B29=1,Assumptions!B17,B19)", inputs["discount_rate"]),
    }
    for row, (label, expression, value) in wacc_links.items():
        text("WACC", row, 1, label)
        formula("WACC", row, 2, expression, value, MONEY if row == 8 else PERCENT)
    _note(
        wb["WACC"],
        23,
        "CAPM Ke = Rf + beta × ERP. WACC = Ke × equity weight + pre-tax borrowing cost × (1 − tax rate) × debt weight. Book debt and market equity may have different observation dates.",
    )

    wc_rows = {row["year"]: row for row in wc.get("forecast", [])}
    labels = {
        6: "Revenue",
        8: "Cost of revenue",
        9: "Operating profit (EBIT)",
        10: "Tax rate",
        11: "NOPAT",
        12: "Depreciation",
        13: "Capital expenditure",
        14: "Operating working capital",
        15: "Change in working capital",
        16: "Free cash flow to firm",
    }
    for row, label in labels.items():
        text("OperatingForecast", row, 1, label)
    for row, label in {
        6: "Revenue",
        7: "Trade receivables",
        8: "Inventory",
        9: "Trade payables",
        10: "Operating NWC",
        11: "Change in operating NWC",
    }.items():
        text("WorkingCapital", row, 1, label)
    for col in range(2, 8):
        letter = get_column_letter(col)
        previous = get_column_letter(col - 1)
        year = model["base_year"] + col - 2
        for sheet in ("OperatingForecast", "WorkingCapital", "DCF"):
            text(sheet, 5, col, f"{year}{'A' if col == 2 else 'E'}")
        if col == 2:
            formula("OperatingForecast", 6, col, "=Assumptions!B7", base["revenue"] / divisor)
            formula("WorkingCapital", 6, col, "=Assumptions!B7", base["revenue"] / divisor)
            for row, key in ((7, "accounts_receivable"), (8, "inventory"), (9, "accounts_payable")):
                raw = base.get(key)
                if wc.get("base_observations"):
                    formula(
                        "WorkingCapital", row, col, f"={ref(key)}/Assumptions!$B$6", raw / divisor
                    )
            if wc.get("base_observations"):
                expr = "=B7+B8-B9"
            else:
                expr = "=IF(Assumptions!$B$33=1,B6*Assumptions!$B$34/365+B6*Assumptions!$B$37*(Assumptions!$B$35-Assumptions!$B$36)/365,B6*Assumptions!$B$16)"
            formula("WorkingCapital", 10, col, expr, wc.get("base_balance", 0) / divisor)
            continue
        record = forecast[col - 3]
        wcr = wc_rows.get(year, {})
        expressions = {
            6: f"={previous}6*(1+Assumptions!{letter}46)",
            8: f"={letter}6*Assumptions!{letter}55",
            9: f"={letter}6*Assumptions!{letter}47",
            10: f"=Assumptions!{letter}48",
            11: f"={letter}9-MAX({letter}9,0)*{letter}10",
            12: f"={letter}6*Assumptions!{letter}49",
            13: f"={letter}6*Assumptions!{letter}50",
            14: f"=WorkingCapital!{letter}10",
            15: f"=WorkingCapital!{letter}11",
            16: f"={letter}11+{letter}12-{letter}13-{letter}15",
        }
        keys = {
            6: "revenue",
            8: "cost_of_revenue",
            9: "ebit",
            10: "tax_rate",
            11: "nopat",
            12: "depreciation",
            13: "capital_expenditure",
            14: "working_capital_balance",
            15: "change_in_working_capital",
            16: "free_cash_flow",
        }
        for row, expr in expressions.items():
            value = inputs["tax_rate"] if row == 10 else record.get(keys[row], 0) / divisor
            formula("OperatingForecast", row, col, expr, value, PERCENT if row == 10 else MONEY)
        formula(
            "WorkingCapital", 6, col, f"=OperatingForecast!{letter}6", record["revenue"] / divisor
        )
        for row, key, driver in ((7, "receivables", 52), (8, "inventory", 53), (9, "payables", 54)):
            basis = f"{letter}6" if row == 7 else f"OperatingForecast!{letter}8"
            formula(
                "WorkingCapital",
                row,
                col,
                f'=IF(Assumptions!$B$33=1,{basis}*Assumptions!{letter}{driver}/365,"")',
                wcr.get(key) / divisor if wcr.get(key) is not None else None,
            )
        formula(
            "WorkingCapital",
            10,
            col,
            f"=IF(Assumptions!$B$33=1,{letter}7+{letter}8-{letter}9,{letter}6*Assumptions!{letter}51)",
            record["working_capital_balance"] / divisor,
        )
        formula(
            "WorkingCapital",
            11,
            col,
            f"={letter}10-{previous}10",
            record["change_in_working_capital"] / divisor,
        )
        for row, of_row, key in (
            (6, 6, "revenue"),
            (7, 9, "ebit"),
            (8, 11, "nopat"),
            (9, 12, "depreciation"),
            (10, 13, "capital_expenditure"),
            (11, 15, "change_in_working_capital"),
            (12, 16, "free_cash_flow"),
        ):
            text("DCF", row, 1, labels.get(of_row, key))
            formula("DCF", row, col, f"=OperatingForecast!{letter}{of_row}", record[key] / divisor)
        text("DCF", 13, 1, "Discount period (actual days / 365.25)")
        formula(
            "DCF",
            13,
            col,
            f"=({letter}17-Assumptions!$B$58)/365.25",
            record["discount_period"],
            "0.0000",
        )
        text("DCF", 14, 1, "Active WACC")
        formula("DCF", 14, col, "=WACC!$B$20", inputs["discount_rate"], PERCENT)
        text("DCF", 15, 1, "Discount factor")
        formula(
            "DCF",
            15,
            col,
            f'=IF({letter}14>0,1/(1+{letter}14)^{letter}13,"")',
            record["discount_factor"],
            '0.0000"x"',
        )
        text("DCF", 16, 1, "PV of FCFF")
        formula("DCF", 16, col, f"={letter}19*{letter}15", record["present_value"] / divisor)
        text("DCF", 17, 1, "Fiscal cash-flow date")
        formula(
            "DCF",
            17,
            col,
            f"=EDATE(Assumptions!$B$59,12*({col}-2))",
            to_excel(_iso_date(record["cash_flow_date"])),
            "yyyy-mm-dd",
        )
        text("DCF", 18, 1, "Valuation cash-flow fraction")
        fraction_expr = (
            "=MAX(0,MIN(1,(C17-Assumptions!$B$58)/(C17-Assumptions!$B$59)))" if col == 3 else "=1"
        )
        formula("DCF", 18, col, fraction_expr, record["valuation_cash_flow_fraction"], PERCENT)
        text("DCF", 19, 1, "FCFF after valuation date")
        formula(
            "DCF", 19, col, f"={letter}12*{letter}18", record["valuation_free_cash_flow"] / divisor
        )
    minority = finite_number(base.get("noncontrolling_interest"))
    if minority is not None and minority < 0:
        minority = None
    # Cache each supported intermediate even when the per-share valuation is
    # withheld. The saved values must agree with the live formulas in Excel.
    pv_fcff = sum(row["present_value"] for row in forecast)
    rate, growth = inputs["discount_rate"], inputs["terminal_growth"]
    terminal_schedule = model.get("terminal") or {}
    terminal_fcff = terminal_schedule.get("free_cash_flow")
    terminal = (
        terminal_fcff / (rate - growth)
        if rate > growth and terminal_fcff is not None and terminal_fcff > 0
        else None
    )
    pv_terminal = (
        terminal / (1 + rate) ** forecast[-1]["discount_period"] if terminal is not None else None
    )
    enterprise = pv_fcff + pv_terminal if pv_terminal is not None else None
    equity = enterprise - inputs["net_debt"] - (minority or 0) if enterprise is not None else None
    quote = finite_number(snapshot.get("current_price"))
    quote = quote if quote is not None and quote > 0 and comparison_available else None
    bridge = {
        21: ("PV of forecast FCFF", "=SUM(C16:G16)", pv_fcff),
        22: (
            "Terminal value",
            '=IF(AND(Assumptions!$B$61=1,WACC!$B$20>B40,B41>0,B42>=0,B42<=1,B44>0),B44/(WACC!$B$20-B40),"")',
            terminal,
        ),
        23: (
            "PV of terminal value",
            '=IF(ISNUMBER(B22),B22/(1+WACC!$B$20)^G13,"")',
            pv_terminal,
        ),
        24: (
            "Enterprise value",
            '=IF(ISNUMBER(B23),B21+B23,"")',
            enterprise,
        ),
        26: ("Net debt / (cash)", "=Assumptions!B19", inputs["net_debt"]),
        27: (
            "Equity value",
            '=IF(ISNUMBER(B24),B24-B26-IF(ISNUMBER(B25),B25,0),"")',
            equity,
        ),
    }
    for row, (label, expr, value) in bridge.items():
        text("DCF", row, 1, label)
        formula("DCF", row, 2, expr, value / divisor if value is not None else None)
    text("DCF", 25, 1, "Non-controlling interests (book proxy)")
    formula(
        "DCF",
        25,
        2,
        f'=IF(AND(ISNUMBER({ref("noncontrolling_interest")}),{ref("noncontrolling_interest")}>=0),{ref("noncontrolling_interest")}/Assumptions!$B$6,"")',
        minority / divisor if minority is not None else None,
    )
    text("DCF", 28, 1, "Illustrative value / share")
    formula(
        "DCF",
        28,
        2,
        '=IF(AND(ISNUMBER(B27),B27>0,Assumptions!$B$20>0),B27*Assumptions!$B$6/Assumptions!$B$20,"")',
        valuation.get("per_share"),
    )
    text("DCF", 29, 1, "Observed quote / share")
    formula(
        "DCF",
        29,
        2,
        '=IF(AND(Assumptions!$B$62=1,ISNUMBER(Sources!B8),Sources!B8>0),Sources!B8,"")',
        quote,
    )
    text("DCF", 30, 1, "Model / market difference")
    formula(
        "DCF",
        30,
        2,
        '=IF(AND(ISNUMBER(B28),ISNUMBER(B29),B29>0),B28/B29-1,"")',
        valuation.get("upside_percent") / 100
        if valuation.get("upside_percent") is not None
        else None,
        PERCENT,
    )
    _note(
        wb["DCF"],
        33,
        "Terminal FCFF = terminal NOPAT × (1 − growth / ROIC); terminal value = terminal FCFF / (WACC − growth). Reported minority interests are a book-value proxy. Blank means adjustment omitted, not verified zero. Invalid cash flow, stale annual coverage or nonpositive equity withholds per-share value.",
    )
    for row, label, expr, key, fmt in (
        (36, "Terminal financial year", "=Assumptions!B8+6", "year", "0"),
        (37, "Terminal revenue", "=G6*(1+B40)", "revenue", MONEY),
        (38, "Terminal EBIT", "=B37*Assumptions!G47", "ebit", MONEY),
        (39, "Terminal NOPAT", "=B38-MAX(B38,0)*Assumptions!G48", "nopat", MONEY),
        (40, "Terminal growth", "=Assumptions!B18", "growth", PERCENT),
        (41, "Terminal ROIC", "=IF(Assumptions!B60=0,WACC!B20,Assumptions!B40)", "roic", PERCENT),
        (42, "Terminal reinvestment rate", '=IF(B41>0,B40/B41,"")', "reinvestment_rate", PERCENT),
        (43, "Terminal reinvestment", '=IF(ISNUMBER(B42),B39*B42,"")', "reinvestment", MONEY),
        (44, "Terminal free cash flow", '=IF(ISNUMBER(B43),B39-B43,"")', "free_cash_flow", MONEY),
    ):
        text("DCF", row, 1, label)
        value = terminal_schedule.get(key)
        formula(
            "DCF",
            row,
            2,
            expr,
            value / divisor if fmt == MONEY and value is not None else value,
            fmt,
        )
    for row, expr, value, fmt in (
        (8, "=DCF!B24", enterprise, MONEY),
        (9, "=DCF!B27", equity, MONEY),
        (12, "=DCF!B28", valuation.get("per_share"), MONEY),
        (13, "=DCF!B29", quote, MONEY),
        (14, "=DCF!B30", valuation.get("upside_percent"), PERCENT),
        (16, "=WACC!B20", inputs["discount_rate"], PERCENT),
        (17, "=Assumptions!B18", inputs["terminal_growth"], PERCENT),
    ):
        cached = (
            value / divisor
            if row in (8, 9) and value is not None
            else value / 100
            if row == 14 and value is not None
            else value
        )
        formula("Summary", row, 2, expr, cached, fmt)
    sensitivity = model.get("sensitivity", {})
    rates = sensitivity.get("discount_rates", [])
    growths = sensitivity.get("terminal_growth_rates", [])
    for j, g in enumerate(growths, 2):
        offset = j - 4
        expr = (
            f"=MAX(0,Assumptions!$B$18+({offset})*MIN(0.01,Assumptions!$B$18/2))"
            if inputs["terminal_growth"]
            else f"=Assumptions!$B$18+{j - 2}*0.01"
        )
        formula("Sensitivity", 7, j, expr, g, PERCENT)
    for i, r in enumerate(rates, 8):
        formula("Sensitivity", i, 1, f"=WACC!$B$20+({i - 10})*MIN(0.01,WACC!$B$20/3)", r, PERCENT)
        for j, _ in enumerate(growths, 2):
            col = get_column_letter(j)
            pv = "+".join(
                f"DCF!{get_column_letter(k)}19/(1+$A{i})^DCF!{get_column_letter(k)}13"
                for k in range(3, 8)
            )
            roic = f"IF(Assumptions!$B$60=0,$A{i},Assumptions!$B$40)"
            terminal_cf = f"(DCF!$G$8*(1+{col}$7)*(1-{col}$7/({roic})))"
            equity = f"({pv}+{terminal_cf}/($A{i}-{col}$7)/(1+$A{i})^DCF!$G$13-DCF!$B$26-IF(ISNUMBER(DCF!$B$25),DCF!$B$25,0))"
            expr = f'=IF(AND(Assumptions!$B$61=1,$A{i}>{col}$7,{col}$7>=0,({roic})>{col}$7,DCF!$G$8>0,Assumptions!$B$20>0),IF({equity}>0,{equity}*Assumptions!$B$6/Assumptions!$B$20,""),"")'
            formula("Sensitivity", i, j, expr, sensitivity.get("values", [])[i - 8][j - 2])
    _note(
        wb["Sensitivity"],
        16,
        "Per-share values, unscaled. Rows change WACC; columns change terminal growth. Every cell recomputes terminal NOPAT, growth / ROIC reinvestment and terminal FCFF. Default ROIC follows row WACC; explicit ROIC remains fixed. Forecast cash flows use the remaining first-year fraction and actual-day discount periods.",
    )
    for sheet in wb:
        sheet.print_area = f"A1:H{max(sheet.max_row, 25)}"
        for row in sheet.iter_rows(min_row=4):
            for cell in row:
                if cell.alignment.wrap_text is not True:
                    cell.alignment = Alignment(
                        vertical="center",
                        horizontal="left" if cell.column == 1 else "right",
                        wrap_text=cell.column == 1,
                    )
                    if cell.column == 1 and isinstance(cell.value, str) and len(cell.value) > 35:
                        sheet.row_dimensions[cell.row].height = max(
                            sheet.row_dimensions[cell.row].height or 15, 30
                        )
        sheet.print_title_rows = "1:5"
    return _cached_bytes(wb, caches)
