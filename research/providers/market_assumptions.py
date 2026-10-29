"""Dated, keyless India discount-rate inputs from RBI and NYU Stern.

The RBI observation is a roughly ten-year INR sovereign yield, not a default-free
quote. Subtracting NYU's rating-based sovereign spread creates an explicitly
labelled estimate. Country ERP is a country proxy, not a company-specific fact.
"""

from __future__ import annotations

import copy
import io
import json
import math
import re
import time
import zipfile
from datetime import UTC, date, datetime
from pathlib import Path
from urllib.parse import urljoin, urlsplit
from uuid import uuid4

import httpx
from bs4 import BeautifulSoup, Comment
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

RBI_URL = "https://www.rbi.org.in/"
NYU_INDEX_URL = "https://pages.stern.nyu.edu/~adamodar/New_Home_Page/datacurrent.html"
_HOSTS = {"www.rbi.org.in", "pages.stern.nyu.edu"}
_HTML_LIMIT = 2 * 1024 * 1024
_XLSX_LIMIT = 8 * 1024 * 1024
_EXPANDED_LIMIT = 40 * 1024 * 1024
_CACHE_VERSION = 1
_MONTH_DATE = re.compile(
    r"\b(January|February|March|April|May|June|July|August|September|October|November|December)"
    r"\s+(\d{1,2}),?\s+(20\d{2})\b",
    re.I,
)


def _today() -> date:
    return datetime.now(UTC).date()


def _valid_url(url: str) -> bool:
    try:
        parsed = urlsplit(url)
        return (
            parsed.scheme == "https"
            and parsed.hostname in _HOSTS
            and parsed.port in (None, 443)
            and parsed.username is None
            and parsed.password is None
            and not parsed.fragment
            and not any(ord(char) < 33 or char == "\\" for char in url)
        )
    except (TypeError, ValueError):
        return False


def _date(text: str) -> date | None:
    match = _MONTH_DATE.search(text)
    if not match:
        return None
    try:
        return datetime.strptime(" ".join(match.groups()), "%B %d %Y").date()
    except ValueError:
        return None


def _fresh(as_of: str, limit: int, today: date) -> bool:
    try:
        return 0 <= (today - date.fromisoformat(as_of)).days <= limit
    except (TypeError, ValueError):
        return False


def _number(value, upper: float = 0.6) -> float | None:
    if isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError, OverflowError):
        return None
    return number if math.isfinite(number) and 0 <= number <= upper else None


def _observation(value=None, as_of=None, source_url=None, label="Unavailable", **extra) -> dict:
    return {"value": value, "as_of": as_of, "source_url": source_url, "label": label, **extra}


def _download(client: httpx.Client, url: str, limit: int, deadline: float) -> bytes:
    """Three fixed-host GETs at most; never follow a redirect or use env proxies."""
    if not _valid_url(url) or time.monotonic() >= deadline:
        raise ValueError("Public market source unavailable")
    remaining = deadline - time.monotonic()
    # Bound each blocking phase as well as checking the overall stream deadline.
    timeout = httpx.Timeout(max(0.05, min(2.0, remaining / 4)))
    with client.stream("GET", url, timeout=timeout) as response:
        if response.status_code != 200:
            raise ValueError("Public market source did not return a usable response")
        length = response.headers.get("content-length")
        if length and (not length.isdigit() or int(length) > limit):
            raise ValueError("Public market source exceeds the download limit")
        chunks, size = [], 0
        for chunk in response.iter_bytes(chunk_size=65536):
            size += len(chunk)
            if size > limit or time.monotonic() >= deadline:
                raise ValueError("Public market source exceeds its bounds")
            chunks.append(chunk)
        return b"".join(chunks)


def _rbi_yield(body: bytes, today: date) -> dict:
    soup = BeautifulSoup(body, "html.parser")
    # RBI retains older dates in comments alongside the current as-of date.
    for comment in soup.find_all(string=lambda item: isinstance(item, Comment)):
        comment.extract()
    heading = next(
        (
            item
            for item in soup.find_all(["h2", "h3", "h4"])
            if item.get_text(" ", strip=True) == "Government Securities Market"
        ),
        None,
    )
    section = heading.find_next_sibling("div") if heading else None
    if section is None:
        raise ValueError("RBI government-security section unavailable")
    dates = []
    for line in section.find_all("div"):
        text = line.get_text(" ", strip=True)
        if re.search(r"#\s*as on", text, re.I):
            found = _date(text)
            if found:
                dates.append(found)
    if len(set(dates)) != 1 or not _fresh(dates[0].isoformat(), 10, today):
        raise ValueError("RBI sovereign-yield date is missing or stale")
    as_of = dates[0]
    candidates = []
    for row in section.find_all("tr"):
        cells = row.find_all(["th", "td"])
        if len(cells) != 2:
            continue
        security, quote = (cell.get_text(" ", strip=True) for cell in cells)
        maturity = re.fullmatch(r"\d+(?:\.\d+)?%\s+GS\s+(20\d{2})", security)
        ytm = re.fullmatch(r":?\s*(\d+(?:\.\d+)?)%\s*#", quote)
        if maturity and ytm:
            years = int(maturity.group(1)) - as_of.year
            value = _number(float(ytm.group(1)) / 100, 0.3)
            if 8 <= years <= 12 and value is not None:
                candidates.append((abs(years - 10), security, value))
    if not candidates:
        raise ValueError("RBI roughly ten-year sovereign yield unavailable")
    candidates.sort()
    if len(candidates) > 1 and candidates[0][0] == candidates[1][0]:
        raise ValueError("RBI roughly ten-year sovereign yield is ambiguous")
    _, security, value = candidates[0]
    return _observation(
        value,
        as_of.isoformat(),
        RBI_URL,
        f"RBI {security} yield; approximately 10-year INR sovereign-yield proxy",
        security=security,
        quote_convention="Published annual sovereign yield; not the coupon rate",
    )


def _latest_workbook(body: bytes, today: date) -> tuple[str, date]:
    soup = BeautifulSoup(body, "html.parser")
    candidates = []
    for link in soup.find_all("a", href=True):
        url = urljoin(NYU_INDEX_URL, link["href"])
        published = _date(link.get_text(" ", strip=True))
        path = urlsplit(url).path
        if (
            _valid_url(url)
            and urlsplit(url).hostname == "pages.stern.nyu.edu"
            and re.fullmatch(r"/~adamodar/pc/datasets/ctryprem[A-Za-z0-9]*\.xlsx", path)
            and published is not None
            and published <= today
        ):
            candidates.append((published, url))
    if not candidates:
        raise ValueError("NYU dated country-premium workbook unavailable")
    published, url = max(candidates)
    if not _fresh(published.isoformat(), 450, today):
        raise ValueError("NYU dated country-premium workbook is stale")
    return url, published


def _key(value) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _country_premiums(body: bytes, source_url: str, published: date) -> tuple[dict, dict, list]:
    if len(body) > _XLSX_LIMIT:
        raise ValueError("NYU workbook exceeds download limit")
    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        entries = archive.infolist()
        if len(entries) > 512 or sum(entry.file_size for entry in entries) > _EXPANDED_LIMIT:
            raise ValueError("NYU workbook exceeds expanded-size limits")
        for entry in entries:
            if entry.flag_bits & 1:
                raise ValueError("Encrypted workbook is unsupported")
            if entry.filename.endswith(".xml"):
                xml = archive.read(entry).upper()
                if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                    raise ValueError("Workbook XML declarations are unsupported")
    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True, keep_links=False)
    warnings = []
    try:
        if "ERPs by country" not in workbook.sheetnames:
            raise ValueError("NYU country-premium worksheet unavailable")
        sheet = workbook["ERPs by country"]
        columns, india = None, []
        for row in sheet.iter_rows(min_row=1, max_row=500, max_col=32, values_only=True):
            keys = {_key(cell): index for index, cell in enumerate(row) if cell is not None}
            if "country" in keys and "totalequityriskpremium" in keys:
                spread = keys.get("ratingbaseddefaultspread", keys.get("adjdefaultspread"))
                if spread is not None:
                    columns = (keys["country"], keys["totalequityriskpremium"], spread)
            if columns and str(row[columns[0]] or "").strip() == "India":
                india.append((_number(row[columns[1]]), _number(row[columns[2]], 0.3)))
        if len(india) != 1 or any(value is None for value in india[0]):
            raise ValueError("NYU India rating-based premium observations unavailable")
        erp, spread = india[0]
        internal_date = sheet["B2"].value
        if isinstance(internal_date, (int, float)) and 30000 <= internal_date <= 60000:
            internal_date = from_excel(internal_date)
        if isinstance(internal_date, (date, datetime)):
            internal_date = (
                internal_date.date() if isinstance(internal_date, datetime) else internal_date
            )
            if abs((internal_date - published).days) > 31:
                warnings.append(
                    "NYU workbook date cell differs from the official dated download link; "
                    "the disclosed as-of date follows the official index publication label."
                )
    finally:
        workbook.close()
    basis = "Publication date on the official NYU Stern dataset index; not a daily market quote"
    return (
        _observation(
            erp,
            published.isoformat(),
            source_url,
            "NYU Stern India total equity risk premium (rating-based country proxy)",
            date_basis=basis,
        ),
        _observation(
            spread,
            published.isoformat(),
            source_url,
            "NYU Stern India rating-based sovereign default spread estimate",
            date_basis=basis,
        ),
        warnings,
    )


def _cache_path(settings) -> Path:
    root = Path(getattr(settings, "artifact_storage_root", ".private/artifacts"))
    return root / "cache" / "market-assumptions" / "india.json"


def _read_cache(settings, today: date) -> dict | None:
    try:
        path = _cache_path(settings)
        if path.stat().st_size > 131072:
            return None
        cached = json.loads(path.read_text(encoding="utf-8"))
        retrieved = datetime.fromisoformat(cached["retrieved_at"])
        age = (datetime.now(UTC) - retrieved).total_seconds()
        if cached.get("cache_version") != _CACHE_VERSION or not 0 <= age <= 86400:
            return None
        if cached.get("status") != "complete":
            return None
        for key, limit in (
            ("risk_free_rate", 10),
            ("sovereign_yield", 10),
            ("equity_risk_premium", 450),
            ("sovereign_default_spread", 450),
        ):
            observation = cached[key]
            if (
                _number(observation["value"]) is None
                or not _fresh(observation["as_of"], limit, today)
                or not _valid_url(observation["source_url"])
            ):
                return None
        cached["cache_hit"] = True
        return copy.deepcopy(cached)
    except (OSError, TypeError, ValueError, KeyError):
        return None


def _write_cache(settings, result: dict) -> None:
    if result["status"] != "complete":
        return
    temporary = None
    try:
        path = _cache_path(settings)
        path.parent.mkdir(parents=True, exist_ok=True)
        temporary = path.with_suffix("." + uuid4().hex + ".tmp")
        temporary.write_text(json.dumps(result, allow_nan=False), encoding="utf-8")
        temporary.replace(path)
    except OSError:
        pass
    finally:
        if temporary is not None:
            try:
                temporary.unlink(missing_ok=True)
            except OSError:
                pass


def fetch_india_market_assumptions(settings=None, refresh: bool = False) -> dict:
    """Return observations or explicit missing values; never manufacture rates.

    Only three bounded public HTTPS requests are made. The 24-hour cache contains
    public rates and citations; it is never served after source-date validation
    fails or as a stale fallback after a refresh failure.
    """
    today = _today()
    if not refresh:
        cached = _read_cache(settings, today)
        if cached is not None:
            return cached
    result = {
        "cache_version": _CACHE_VERSION,
        "retrieved_at": datetime.now(UTC).isoformat(),
        "risk_free_rate": _observation(),
        "sovereign_yield": _observation(),
        "equity_risk_premium": _observation(),
        "sovereign_default_spread": _observation(),
        "warnings": [],
        "citations": [],
        "status": "unavailable",
        "cache_hit": False,
    }
    # Reserve time inside the 15-second fetch budget for the final bounded read
    # and workbook parsing; each blocking network phase is at most two seconds.
    deadline = time.monotonic() + 12
    with httpx.Client(
        follow_redirects=False,
        trust_env=False,
        headers={"User-Agent": "Astra-Research/1.0 public-market-data"},
    ) as client:
        try:
            result["sovereign_yield"] = _rbi_yield(
                _download(client, RBI_URL, _HTML_LIMIT, deadline), today
            )
        except Exception:
            result["warnings"].append(
                "A fresh, dated RBI approximately 10-year INR sovereign yield could not be verified; "
                "supply an explicit risk-free-rate assumption."
            )
        try:
            url, published = _latest_workbook(
                _download(client, NYU_INDEX_URL, _HTML_LIMIT, deadline), today
            )
            erp, spread, warnings = _country_premiums(
                _download(client, url, _XLSX_LIMIT, deadline), url, published
            )
            result["equity_risk_premium"], result["sovereign_default_spread"] = erp, spread
            result["warnings"].extend(warnings)
        except Exception:
            result["warnings"].append(
                "A fresh, dated NYU Stern India equity risk premium and sovereign spread could not "
                "be verified; supply an explicit equity-risk-premium assumption."
            )
    sovereign, spread = result["sovereign_yield"], result["sovereign_default_spread"]
    if sovereign["value"] is not None:
        if spread["value"] is not None and sovereign["value"] > spread["value"]:
            result["risk_free_rate"] = _observation(
                sovereign["value"] - spread["value"],
                sovereign["as_of"],
                RBI_URL,
                "Estimated INR default-free proxy: RBI sovereign yield less NYU rating-based default spread",
                adjustment_source_url=spread["source_url"],
                adjustment_as_of=spread["as_of"],
                basis="Derived estimate using observations of different dates; not a quoted default-free rate",
            )
            result["warnings"].append(
                "Risk-free input is an adjusted estimate using the disclosed RBI yield and NYU "
                "sovereign-spread dates; India ERP is a country proxy and may differ from company exposure."
            )
        elif spread["value"] is None:
            result["risk_free_rate"] = {
                **sovereign,
                "label": "Unadjusted INR sovereign-yield proxy; default spread unavailable",
            }
            result["warnings"].append(
                "The INR government-bond yield proxy is unadjusted for sovereign default risk; "
                "it is not a verified default-free rate."
            )
        else:
            result["warnings"].append(
                "The sovereign spread does not produce a positive adjusted rate; enter an explicit risk-free assumption."
            )
    complete = all(
        result[key]["value"] is not None
        for key in ("risk_free_rate", "equity_risk_premium", "sovereign_default_spread")
    )
    result["status"] = (
        "complete"
        if complete
        else (
            "partial"
            if any(
                result[key]["value"] is not None
                for key in ("sovereign_yield", "equity_risk_premium")
            )
            else "unavailable"
        )
    )
    for key in ("sovereign_yield", "equity_risk_premium"):
        observation = result[key]
        if observation["value"] is not None:
            result["citations"].append(
                {
                    "title": observation["label"],
                    "uri": observation["source_url"],
                    "as_of": observation["as_of"],
                    "retrieved_label": result["retrieved_at"],
                    "provider": "rbi" if key == "sovereign_yield" else "nyu-stern",
                }
            )
    _write_cache(settings, result)
    return result
