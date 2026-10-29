import io
import zipfile
from datetime import date

import pytest
from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils import get_column_letter, range_boundaries

from research.data.stock_model import build_stock_model
from research.reporting.excel_model import build_stock_workbook
from tests.test_stock_model import stock_observation as observation_fixture


def test_adjusted_risk_free_rate_keeps_yield_and_default_spread_linked(stock_observation):
    stock_observation["market_assumptions"] = {
        "risk_free_rate": {
            "value": 0.05,
            "as_of": "2025-03-31",
            "source_url": "https://www.rbi.org.in/",
            "adjustment_source_url": "https://pages.stern.nyu.edu/",
        },
        "sovereign_yield": {
            "value": 0.07,
            "as_of": "2025-03-31",
            "source_url": "https://www.rbi.org.in/",
        },
        "sovereign_default_spread": {
            "value": 0.02,
            "as_of": "2025-03-31",
            "source_url": "https://pages.stern.nyu.edu/",
        },
        "equity_risk_premium": {
            "value": 0.06,
            "as_of": "2025-03-31",
            "source_url": "https://pages.stern.nyu.edu/",
        },
    }
    model = build_stock_model(stock_observation)
    data = build_stock_workbook(model)
    workbook = load_workbook(io.BytesIO(data))
    source = workbook["Sources"]
    rows = {source.cell(row, 1).value: row for row in range(1, source.max_row + 1)}
    assert (
        workbook["Assumptions"]["B24"].value
        == f"=Sources!B{rows['Sovereign Yield']}-Sources!B{rows['Sovereign Default Spread']}"
    )
    cached = load_workbook(io.BytesIO(data), data_only=True)
    assert cached["Assumptions"]["B24"].value == pytest.approx(0.05)


@pytest.fixture
def stock_observation():
    return observation_fixture.__wrapped__()


def test_workbook_contains_editable_formulas_and_matching_cached_outputs(stock_observation):
    model = build_stock_model(stock_observation)
    content = build_stock_workbook(model)
    formulas = load_workbook(io.BytesIO(content), data_only=False)
    values = load_workbook(io.BytesIO(content), data_only=True)
    assert formulas.sheetnames == [
        "Summary",
        "Assumptions",
        "Historical",
        "OperatingForecast",
        "WorkingCapital",
        "WACC",
        "DCF",
        "Sensitivity",
        "Sources",
    ]
    assert formulas["DCF"]["C6"].value == "=OperatingForecast!C6"
    assert formulas["OperatingForecast"]["C6"].value == "=B6*(1+Assumptions!C46)"
    assert formulas["DCF"]["B28"].value.startswith("=IF(AND(")
    assert "Assumptions!$B$6/Assumptions!$B$20" in formulas["DCF"]["B28"].value
    assert values["Summary"]["B12"].value == pytest.approx(model["valuation"]["per_share"])
    assert values["DCF"]["G6"].value == pytest.approx(
        model["forecast"][-1]["revenue"] / model["unit"]["divisor"]
    )
    assert values["Assumptions"]["B19"].value == 50
    assert values["Assumptions"]["B20"].value == 100_000_000
    assert values["Sensitivity"]["D10"].value == pytest.approx(model["valuation"]["per_share"])
    assert formulas.calculation.fullCalcOnLoad and formulas.calculation.forceFullCalc
    assert formulas["Assumptions"].data_validations.count == 25
    assert formulas["Assumptions"]["B29"].number_format == "0"
    assert formulas["Assumptions"]["B33"].number_format == "0"
    assert "absolute source amounts" in formulas["Sources"]["A3"].value
    assert formulas["Historical"]["B7"].value.startswith("=Sources!")


def test_workbook_scaling_changes_money_display_but_not_per_share(stock_observation):
    results = {}
    for unit in ("crore", "million", "lakh", "absolute"):
        model = build_stock_model(stock_observation, display_unit=unit)
        workbook = load_workbook(io.BytesIO(build_stock_workbook(model)), data_only=True)
        results[unit] = workbook["Summary"]["B12"].value
        assert workbook["Historical"]["C7"].value == pytest.approx(
            12_000_000_000 / model["unit"]["divisor"]
        )
        assert workbook["DCF"]["B26"].value == pytest.approx(500_000_000 / model["unit"]["divisor"])
    assert len(set(results.values())) == 1


def test_withheld_bank_workbook_keeps_actuals_and_avoids_fake_valuation(stock_observation):
    stock_observation["sector"] = "Banking"
    workbook = load_workbook(
        io.BytesIO(build_stock_workbook(build_stock_model(stock_observation))), data_only=False
    )
    assert "unsuitable" in workbook["DCF"]["A6"].value
    assert workbook["Summary"]["B12"].value is None
    assert workbook["Historical"]["C7"].value.startswith("=Sources!")


def test_provider_strings_are_excel_text_not_formulas(stock_observation):
    stock_observation["name"] = '=HYPERLINK("https://example.org","fake")'
    stock_observation["citations"][0]["title"] = "=1+1"
    content = build_stock_workbook(build_stock_model(stock_observation))
    workbook = load_workbook(io.BytesIO(content), data_only=False)
    assert workbook["Sources"]["B4"].value.startswith("=HYPERLINK")
    assert workbook["Sources"]["B4"].data_type == "s"
    assert any(
        cell.data_type == "s" and cell.value == "=1+1"
        for row in workbook["Sources"]
        for cell in row
    )
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        assert not any(
            "externalLink" in name or "vbaProject" in name for name in archive.namelist()
        )


def test_explicit_overrides_are_editable_and_retained(stock_observation):
    model = build_stock_model(stock_observation, {"revenue_growth": 0.08, "net_debt": -100_000_000})
    workbook = load_workbook(io.BytesIO(build_stock_workbook(model)), data_only=False)
    assert workbook["Assumptions"]["B11"].value == 0.08
    assert workbook["Assumptions"]["B19"].value == -10
    assert workbook["Assumptions"]["B11"].font.color.rgb == "001768B4"


def test_missing_price_remains_blank_without_removing_supported_dcf(stock_observation):
    stock_observation["current_price"] = None
    model = build_stock_model(stock_observation)
    workbook = load_workbook(io.BytesIO(build_stock_workbook(model)), data_only=True)
    assert model["valuation"]["per_share"] > 0
    assert workbook["Summary"]["B11"].value is None
    assert workbook["Summary"]["B13"].value is None


def _formula_dependencies(workbook):
    """Follow the actual workbook references, including cross-sheet ranges."""
    graph = {}
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if cell.data_type != "f":
                    continue
                references = set()
                for token in Tokenizer(cell.value).items:
                    if token.type != "OPERAND" or token.subtype != "RANGE":
                        continue
                    location = token.value.replace("$", "")
                    owner, address = (
                        location.rsplit("!", 1) if "!" in location else (sheet.title, location)
                    )
                    owner = owner.strip("'")
                    assert owner in workbook.sheetnames, f"Unknown sheet: {owner}"
                    left, top, right, bottom = range_boundaries(address)
                    for column in range(left, right + 1):
                        for line in range(top, bottom + 1):
                            references.add((owner, f"{get_column_letter(column)}{line}"))
                graph[sheet.title, cell.coordinate] = references
    return graph


def test_forecast_wacc_working_capital_and_all_yearly_inputs_reach_valuation(stock_observation):
    workbook = load_workbook(io.BytesIO(build_stock_workbook(build_stock_model(stock_observation))))
    graph = _formula_dependencies(workbook)
    visited, stack = set(), set()

    def visit(node):
        assert node not in stack, f"Circular formula reference: {node}"
        if node in visited:
            return
        stack.add(node)
        for dependency in graph.get(node, ()):
            visit(dependency)
        stack.remove(node)
        visited.add(node)

    for node in graph:
        visit(node)

    def ancestors(node):
        result = set()
        pending = [node]
        while pending:
            for dependency in graph.get(pending.pop(), ()):
                if dependency not in result:
                    result.add(dependency)
                    pending.append(dependency)
        return result

    linked = ancestors(("Summary", "B12"))
    for row in (
        11,
        12,
        13,
        14,
        15,
        16,
        17,
        18,
        19,
        20,
        24,
        25,
        26,
        27,
        28,
        29,
        33,
        34,
        35,
        36,
        37,
        39,
        40,
        58,
        59,
        60,
        61,
    ):
        assert ("Assumptions", f"B{row}") in linked
    for column in ("C", "D", "E", "F", "G"):
        for row in range(46, 56):
            assert ("Assumptions", f"{column}{row}") in linked
    for flag in ("B29", "B33", "B60", "B62"):
        validation = next(
            rule
            for rule in workbook["Assumptions"].data_validations.dataValidation
            if flag in rule.sqref
        )
        assert validation.type == "whole"


@pytest.mark.parametrize("terminal_roic", [None, 0.18])
def test_workbook_terminal_and_stub_cashflows_match_independent_math(
    stock_observation, terminal_roic
):
    stock_observation["price_as_of"] = "2025-10-15"
    overrides = {
        "revenue_growth": 0.2,
        "revenue_growth_end": 0.04,
        "discount_rate": 0.12,
        "terminal_growth": 0.03,
    }
    if terminal_roic is not None:
        overrides["terminal_roic"] = terminal_roic
    model = build_stock_model(stock_observation, overrides)
    content = build_stock_workbook(model)
    formulas = load_workbook(io.BytesIO(content))
    values = load_workbook(io.BytesIO(content), data_only=True)
    divisor = 10_000_000
    date_of_value = date(2025, 10, 15)
    previous_revenue = 12_000_000_000
    previous_wc = previous_revenue * 0.02
    cashflows, periods = [], []
    for step in range(1, 6):
        growth = 0.2 + (0.04 - 0.2) * (step - 1) / 4
        revenue = previous_revenue * (1 + growth)
        nopat = revenue * 0.2 * 0.75
        working_capital = revenue * 0.02
        fcff = nopat + revenue * 0.03 - revenue * 0.05 - (working_capital - previous_wc)
        cash_date = date(2025 + step, 3, 31)
        period = (cash_date - date_of_value).days / 365.25
        fraction = (cash_date - date_of_value).days / 365 if step == 1 else 1
        cashflows.append(fcff * fraction)
        periods.append(period)
        column = get_column_letter(step + 2)
        assert values["DCF"][f"{column}12"].value == pytest.approx(fcff / divisor)
        assert values["DCF"][f"{column}18"].value == pytest.approx(fraction)
        assert values["DCF"][f"{column}13"].value == pytest.approx(period)
        assert values["DCF"][f"{column}19"].value == pytest.approx(fcff * fraction / divisor)
        assert values["Assumptions"][f"{column}46"].value == pytest.approx(growth)
        previous_revenue, previous_wc = revenue, working_capital
    terminal_nopat = previous_revenue * 1.03 * 0.2 * 0.75
    roic = terminal_roic or 0.12
    terminal_fcff = terminal_nopat * (1 - 0.03 / roic)
    pv_explicit = sum(cf / 1.12**period for cf, period in zip(cashflows, periods, strict=True))
    pv_terminal = terminal_fcff / (0.12 - 0.03) / 1.12 ** periods[-1]
    expected = (pv_explicit + pv_terminal - 500_000_000) / 100_000_000
    assert values["DCF"]["B42"].value == pytest.approx(0.03 / roic)
    assert values["DCF"]["B44"].value == pytest.approx(terminal_fcff / divisor)
    assert values["DCF"]["B23"].value == pytest.approx(pv_terminal / divisor)
    assert values["Summary"]["B12"].value == pytest.approx(expected)
    # Off-centre sensitivity must rebuild growth-supported reinvestment, and
    # use the row WACC for default ROIC rather than retaining base terminal CF.
    sensitivity_rate = values["Sensitivity"]["A8"].value
    sensitivity_growth = values["Sensitivity"]["B7"].value
    sensitivity_roic = terminal_roic or sensitivity_rate
    sensitivity_cf = (
        previous_revenue
        * (1 + sensitivity_growth)
        * 0.2
        * 0.75
        * (1 - sensitivity_growth / sensitivity_roic)
    )
    sensitivity_pv = sum(
        cf / (1 + sensitivity_rate) ** period for cf, period in zip(cashflows, periods, strict=True)
    )
    sensitivity_terminal = (
        sensitivity_cf
        / (sensitivity_rate - sensitivity_growth)
        / (1 + sensitivity_rate) ** periods[-1]
    )
    assert values["Sensitivity"]["B8"].value == pytest.approx(
        (sensitivity_pv + sensitivity_terminal - 500_000_000) / 100_000_000
    )
    assert formulas["DCF"]["C16"].value == "=C19*C15"
    assert "DCF!$G$13" in formulas["Sensitivity"]["B8"].value
    assert "Assumptions!$B$60=0,$A8,Assumptions!$B$40" in formulas["Sensitivity"]["B8"].value
    assert formulas["Assumptions"]["B40"].value == (
        terminal_roic if terminal_roic is not None else "=WACC!B20"
    )


def test_normalized_inputs_link_matching_valid_periods_instead_of_latest_outlier(stock_observation):
    stock_observation["fiscal_period_ends"] = [f"{year}-03-31" for year in range(2022, 2026)]
    stock_observation["financial_history"] = {
        "years": [2022, 2023, 2024, 2025],
        "revenue": [100e6, 120e6, 140e6, 160e6],
        "ebit": [1e6, 24e6, 28e6, 80e6],
        "depreciation": [3e6, 3.6e6, None, 8e6],
        "capital_expenditure": [5e6, 6e6, 7e6, 32e6],
        "total_debt": [10e6] * 4,
        "cash_and_equivalents": [5e6] * 4,
        "pretax_income": [10e6, 10e6, -1e6, 40e6],
        "income_tax": [1e6, 2e6, 3e6, 20e6],
    }
    model = build_stock_model(stock_observation)
    content = build_stock_workbook(model)
    formulas = load_workbook(io.BytesIO(content))
    values = load_workbook(io.BytesIO(content), data_only=True)
    assert values["Assumptions"]["B11"].value == pytest.approx(1.6 ** (1 / 3) - 1)
    assert values["Assumptions"]["B12"].value == pytest.approx(0.2)
    assert values["Assumptions"]["B13"].value == pytest.approx(0.2)
    assert values["Assumptions"]["B14"].value == pytest.approx(0.03)
    assert values["Assumptions"]["B15"].value == pytest.approx(0.05)
    graph = _formula_dependencies(formulas)
    assert ("Sources", "B15") not in graph["Assumptions", "B12"]
    assert ("Sources", "C15") in graph["Assumptions", "B12"]
    assert ("Sources", "D20") not in graph["Assumptions", "B14"]
    assert ("Sources", "B20") in graph["Assumptions", "B14"]
    assert "MEDIAN" in formulas["Assumptions"]["B12"].value
    assert "MIN(0.15" not in formulas["Assumptions"]["B11"].value


def test_undated_quote_is_withheld_by_formula_and_cache(stock_observation):
    stock_observation["price_as_of"] = stock_observation["retrieved_at"] = None
    content = build_stock_workbook(build_stock_model(stock_observation))
    formulas = load_workbook(io.BytesIO(content))
    values = load_workbook(io.BytesIO(content), data_only=True)
    assert values["Assumptions"]["B62"].value == 0
    assert values["Summary"]["B13"].value is None
    assert "Assumptions!$B$62=1" in formulas["DCF"]["B29"].value


def test_stale_annual_coverage_keeps_actuals_and_withholds_workbook_value(stock_observation):
    stock_observation["price_as_of"] = "2026-10-29"
    content = build_stock_workbook(build_stock_model(stock_observation))
    workbook = load_workbook(io.BytesIO(content))
    assert "stale annual coverage" in workbook["DCF"]["A6"].value
    assert workbook["Summary"]["B12"].value is None
    assert workbook["Historical"]["C7"].value.startswith("=Sources!")


@pytest.mark.parametrize("overrides", [{"net_debt": 1_000_000_000_000}, {"ebit_margin": -0.2}])
def test_withheld_valuation_preserves_editable_sensitivity_and_matching_intermediates(
    stock_observation, overrides
):
    model = build_stock_model(stock_observation, overrides)
    assert model["valuation"] is None and len(model["forecast"]) == 5
    content = build_stock_workbook(model)
    formulas = load_workbook(io.BytesIO(content))
    values = load_workbook(io.BytesIO(content), data_only=True)
    assert formulas["Sensitivity"]["D10"].data_type == "f"
    assert values["Summary"]["B12"].value is None
    assert values["DCF"]["B21"].value == pytest.approx(
        sum(row["present_value"] for row in model["forecast"]) / model["unit"]["divisor"]
    )
    if "net_debt" in overrides:
        assert values["DCF"]["B24"].value > 0
        assert values["DCF"]["B27"].value < 0
        assert values["Summary"]["B9"].value == values["DCF"]["B27"].value
    else:
        assert values["DCF"]["B22"].value is None


def test_export_redacts_private_locations_credentials_and_url_parameters(stock_observation):
    stock_observation["name"] = r"Company C:\Users\secret-owner\private-report.pdf"
    stock_observation["citations"] = [
        {
            "title": "Public filing",
            "uri": "https://example.org/filing?api_key=private-token#fragment",
        },
        {"title": "Private source", "uri": "https://username:password@example.org/secret"},
        {"title": "Local source", "uri": "http://localhost:3000/private"},
        {"title": "Private IP", "uri": "https://10.1.2.3/private"},
        {"title": "Uploaded", "uri": "file:///C:/Users/secret-owner/private-report.pdf"},
    ]
    content = build_stock_workbook(build_stock_model(stock_observation))
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        text = " ".join(
            archive.read(name).decode() for name in archive.namelist() if name.endswith(".xml")
        )
    for private in (
        "secret-owner",
        "api_key",
        "private-token",
        "username:password",
        "localhost",
        "10.1.2.3",
        "file://",
    ):
        assert private not in text
    assert "https://example.org/filing" in text
    workbook = load_workbook(io.BytesIO(content), data_only=True)
    assert workbook["Sources"]["C14"].value == 12_000_000_000


def test_numeric_source_fields_cannot_inject_excel_formulas(stock_observation):
    stock_observation["current_price"] = '=HYPERLINK("https://example.org","execute")'
    stock_observation["market_cap"] = "=1+1"
    stock_observation["beta"] = "=1+1"
    model = build_stock_model(stock_observation)
    workbook = load_workbook(io.BytesIO(build_stock_workbook(model)))
    for address in ("B8", "B9", "B10"):
        assert workbook["Sources"][address].data_type != "f"
        assert workbook["Sources"][address].value is None
