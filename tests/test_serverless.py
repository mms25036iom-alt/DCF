import io
from copy import deepcopy

import pytest
from cryptography.fernet import Fernet
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from test_stock_model import stock_observation  # noqa: F401

from services.api.app.core.config import Settings
from services.serverless import app as api

OWNER = "a" * 64
HEADERS = {"Authorization": "Bearer " + "t" * 48, "X-workspace-owner": OWNER}


@pytest.fixture
def client(monkeypatch, tmp_path, stock_observation):  # noqa: F811
    settings = Settings(
        _env_file=None,
        internal_api_token="t" * 48,
        snapshot_encryption_key=Fernet.generate_key().decode(),
        artifact_storage_root=tmp_path,
    )
    monkeypatch.setattr(api, "get_settings", lambda: settings)
    calls = []

    def fetch(_self, ticker, refresh=False):
        calls.append((ticker, refresh))
        result = deepcopy(stock_observation)
        result["ticker"] = api._normal_ticker(ticker)
        return result

    monkeypatch.setattr(api.YahooFinanceProvider, "fetch_company", fetch)
    monkeypatch.setattr(
        api,
        "fetch_india_market_assumptions",
        lambda *_args, **_kwargs: {
            "risk_free_rate": {
                "value": 0.065,
                "as_of": "2025-03-31",
                "source_url": "https://www.rbi.org.in/",
                "label": "Controlled test observation",
            },
            "equity_risk_premium": {
                "value": 0.06,
                "as_of": "2025-03-31",
                "source_url": "https://pages.stern.nyu.edu/",
                "label": "Controlled test estimate",
            },
            "warnings": [],
            "citations": [],
        },
    )
    with TestClient(api.app, raise_server_exceptions=False) as http:
        yield http, calls, settings


def test_private_api_fails_closed_and_no_legacy_routes(client):
    http, _, _ = client
    assert http.get("/health").json() == {"status": "ok"}
    assert http.get("/v1/config").status_code == 401
    assert (
        http.get("/v1/config", headers={**HEADERS, "Authorization": "Bearer wrong"}).status_code
        == 401
    )
    assert http.get("/v1/config", headers=HEADERS).json() == {
        "app_mode": "LIVE",
        "live_ready": True,
        "model_provider": "none",
        "ai_available": False,
    }
    for path in ("/v1/settings", "/v1/research-runs", "/docs", "/openapi.json", "/v1/sources"):
        assert http.get(path, headers=HEADERS).status_code == 404


def test_snapshot_is_encrypted_owner_bound_and_exports_reuse_observations(client):
    http, calls, settings = client
    response = http.post("/v1/stocks/model", headers=HEADERS, json={"ticker": "REAL.NS"})
    assert response.status_code == 200, response.text
    model = response.json()
    token = model["snapshot_id"]
    assert "REAL.NS" not in token and "Verified Industrial" not in token
    payload = {"ticker": "REAL.NS", "snapshot_id": token, "assumptions": {"revenue_growth": 0.08}}
    recalculated = http.post("/v1/stocks/model", headers=HEADERS, json=payload).json()
    assert recalculated["assumptions"]["revenue_growth"] == 0.08
    assert len(calls) == 1
    wrong_owner = http.post(
        "/v1/stocks/model", headers={**HEADERS, "X-workspace-owner": "b" * 64}, json=payload
    )
    assert wrong_owner.status_code == 410
    assert (
        http.post(
            "/v1/stocks/model", headers=HEADERS, json={**payload, "ticker": "OTHER.NS"}
        ).status_code
        == 410
    )
    encrypted = (
        Fernet(settings.snapshot_encryption_key.encode())
        .encrypt_at_time(b"{}", current_time=1)
        .decode()
    )
    assert (
        http.post(
            "/v1/stocks/model", headers=HEADERS, json={**payload, "snapshot_id": encrypted}
        ).status_code
        == 410
    )
    xlsx = http.post("/v1/stocks/model/export", headers=HEADERS, json=payload)
    assert xlsx.status_code == 200, xlsx.text[:300]
    workbook = load_workbook(io.BytesIO(xlsx.content), data_only=True)
    assert workbook["Summary"]["B12"].value == pytest.approx(recalculated["valuation"]["per_share"])
    assert len(workbook.sheetnames) == 9
    pdf = http.post(
        "/v1/equity/report",
        headers=HEADERS,
        json={**payload, "research_brief": "Evaluate capital intensity"},
    )
    assert pdf.status_code == 200, pdf.text[:300]
    pages = PdfReader(io.BytesIO(pdf.content)).pages
    text = "\n".join(page.extract_text() for page in pages)
    assert len(pages) >= 5 and "Evaluate capital intensity" in text
    assert "8.00%" in text and "WACC" in text and "INR crore" in text
    assert len(calls) == 1
    assert token not in text and settings.internal_api_token not in text
    assert "no-store" in pdf.headers["cache-control"]


def test_bounded_bodies_validation_and_errors_do_not_echo_input(client):
    http, calls, _ = client
    assert http.post("/v1/stocks/model", headers=HEADERS, content="x").status_code == 415
    huge = http.post(
        "/v1/stocks/model",
        headers={**HEADERS, "Content-Type": "application/json"},
        content=b"x" * (api.MAX_BODY + 1),
    )
    assert huge.status_code == 413
    invalid = http.post(
        "/v1/stocks/model",
        headers=HEADERS,
        json={"ticker": "../private-secret", "apikey": "do-not-echo"},
    )
    assert invalid.status_code == 422
    assert "do-not-echo" not in invalid.text and "../private-secret" not in invalid.text
    assert not calls


def test_missing_credentials_return_safe_configuration_error(client, monkeypatch):
    http, _, _ = client
    monkeypatch.setattr(api, "get_settings", lambda: Settings(_env_file=None))
    response = http.get("/v1/config", headers=HEADERS)
    assert response.status_code == 503
    assert "Server credentials" in response.text


def test_unsupported_statement_currency_has_a_precise_safe_message(client, monkeypatch):
    http, _, _ = client

    def fail(*_args, **_kwargs):
        raise api.YahooInvalidDataError("Quote currency INR differs from statement currency USD; no FX conversion is configured.")

    monkeypatch.setattr(api.YahooFinanceProvider, "fetch_company", fail)
    response = http.post("/v1/stocks/model", headers=HEADERS, json={"ticker": "INFY.NS"})
    assert response.status_code == 422
    assert "no FX conversion" in response.text


def test_ai_report_receives_compact_financial_evidence_without_credentials(client, monkeypatch):
    http, _, settings = client
    settings.model_router_provider = "groq"
    settings.groq_api_key = "private-test-server-key"
    seen = []

    def generate(_self, _role, _prompt, context, *, qualitative_only=False):
        assert qualitative_only is True
        seen.append(context)
        assert context["scenario_valuation"]["per_share"] > 0
        assert context["metrics"]["terminal"]["reinvestment_rate"] > 0
        assert "snapshot_id" not in str(context)
        assert settings.groq_api_key not in str(context)
        assert settings.internal_api_token not in str(context)
        return "Source-grounded review of the disclosed operating and valuation assumptions."

    monkeypatch.setattr(api.ResearchModelClient, "generate", generate)
    response = http.post(
        "/v1/equity/report", headers=HEADERS, json={"ticker": "REAL.NS", "include_ai": True}
    )
    assert response.status_code == 200
    text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(response.content)).pages)
    assert seen and "Source-grounded review" in text
    assert settings.groq_api_key not in text


def test_ai_numerical_claims_are_withheld_but_calculated_report_remains(client, monkeypatch):
    http, _, settings = client
    settings.model_router_provider = "groq"
    settings.groq_api_key = "private-test-server-key"

    def fail(*_args, **_kwargs):
        raise api.ModelUnavailableError("AI commentary included numerical claims and was withheld.")

    monkeypatch.setattr(api.ResearchModelClient, "generate", fail)
    response = http.post(
        "/v1/equity/report", headers=HEADERS, json={"ticker": "REAL.NS", "include_ai": True}
    )
    assert response.status_code == 200
    text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(response.content)).pages)
    assert "AI commentary included numerical claims and was withheld" in text
    assert "Research synthesis" not in text
    assert "Terminal" in text
