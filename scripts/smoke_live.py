"""Exercise a running app with actual Yahoo observations and matching XLSX/PDF downloads."""

import argparse
import io
import json
import os
import re
from pathlib import Path
from urllib.parse import urlsplit

import httpx
from openpyxl import load_workbook
from pypdf import PdfReader


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-url", default="http://127.0.0.1:3000")
    parser.add_argument("--ticker", default="RELIANCE.NS")
    parser.add_argument("--refresh", action="store_true")
    args = parser.parse_args()
    base = args.base_url.rstrip("/")
    parsed = urlsplit(base)
    if parsed.scheme != "https" and not (
        parsed.scheme == "http" and parsed.hostname in {"127.0.0.1", "localhost"}
    ):
        raise ValueError("Use HTTPS, or a local loopback HTTP instance.")
    with httpx.Client(
        base_url=base, timeout=180, trust_env=False, headers={"Origin": base}
    ) as client:
        password = os.environ.get("APP_ACCESS_PASSWORD")
        if password:
            response = client.post("/api/auth/login", data={"password": password})
            assert response.status_code == 303, "Owner sign-in did not succeed."
        response = client.get("/api/v1/config")
        response.raise_for_status()
        assert response.json()["app_mode"] == "LIVE"
        response = client.post(
            "/api/v1/stocks/model", json={"ticker": args.ticker, "refresh": args.refresh}
        )
        response.raise_for_status()
        original = response.json()
        request = {
            "ticker": original["ticker"],
            "snapshot_id": original["snapshot_id"],
            "display_unit": "auto",
            "assumptions": {"revenue_growth": 0.08},
        }
        recalculated = client.post("/api/v1/stocks/model", json=request)
        recalculated.raise_for_status()
        model = recalculated.json()
        xlsx = client.post("/api/v1/stocks/model/export", json=request)
        xlsx.raise_for_status()
        pdf = client.post(
            "/api/v1/equity/report",
            json={
                **request,
                "research_brief": "Review operating performance, reinvestment needs and valuation sensitivity.",
                "include_ai": False,
            },
        )
        pdf.raise_for_status()
        workbook = load_workbook(io.BytesIO(xlsx.content), data_only=True)
        assert len(workbook.sheetnames) == 9
        if model["valuation"]:
            assert abs(workbook["Summary"]["B12"].value - model["valuation"]["per_share"]) < 1e-8
        reader = PdfReader(io.BytesIO(pdf.content))
        text = "\n".join(page.extract_text() for page in reader.pages)
        assert model["ticker"] in text and "WACC" in text and "8.00%" in text
        assert original["snapshot_id"] not in text
        root = Path("output/verified-stock")
        root.mkdir(parents=True, exist_ok=True)
        name = re.sub(r"[^a-z0-9_-]", "-", model["ticker"].lower())
        (root / f"{name}-dcf.xlsx").write_bytes(xlsx.content)
        (root / f"{name}-equity-research.pdf").write_bytes(pdf.content)
        private = Path(".private/runtime/latest-model.json")
        private.parent.mkdir(parents=True, exist_ok=True)
        private.write_text(json.dumps(model, ensure_ascii=False), encoding="utf-8")
        print(
            json.dumps(
                {
                    "ticker": model["ticker"],
                    "company": model["company"],
                    "unit": model["unit"]["amount_label"],
                    "sheets": workbook.sheetnames,
                    "pdf_pages": len(reader.pages),
                    "scenario_per_share": (model.get("valuation") or {}).get("per_share"),
                    "output_directory": str(root.resolve()),
                },
                indent=2,
            )
        )


if __name__ == "__main__":
    main()
